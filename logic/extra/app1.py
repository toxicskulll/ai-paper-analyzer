import os
import time
import streamlit as st
import pandas as pd
import fitz  # PyMuPDF
from docx import Document
from fpdf import FPDF
import subprocess
from st_aggrid import AgGrid
from st_aggrid.grid_options_builder import GridOptionsBuilder
import tiktoken  # Tokenizer for chunking
import re
import logging

# ------------------- Configuration -------------------
EXTRACTED_DIR = "extracted_data"
SUMMARY_DIR = "summaries"
PDF_DIR = "pdf_summaries"
LOG_FILE = "summarizer_debug.log"

os.makedirs(EXTRACTED_DIR, exist_ok=True)
os.makedirs(SUMMARY_DIR, exist_ok=True)
os.makedirs(PDF_DIR, exist_ok=True)

logging.basicConfig(filename=LOG_FILE, level=logging.DEBUG)

OLLAMA_PATH = r"C:\\Users\\aadis\\AppData\\Local\\Programs\\Ollama\\ollama.exe"
MODEL_NAME = "mistral"

MAX_TOKENS_PER_CHUNK = 800
SUMMARY_RETRIES = 3
RETRY_DELAY_SECONDS = 2

# ------------------- UI Setup -------------------
st.set_page_config(page_title="📄 AI Paper Analyzer", layout="wide")
st.title("📄 AI Research Paper Analyzer")

with st.sidebar:
    depth = st.radio("Summary Style", ["bullet-point", "paragraph", "detailed"], index=0)
    show_chunks = st.checkbox("🔍 Show individual chunk summaries")
    show_logs = st.checkbox("🛠 Show Debug Logs")

# ------------------- Helpers -------------------
def clean_text(text: str) -> str:
    return re.sub(r"[^\x00-\x7F]+", " ", text)

def extract_text_from_file(file_path: str) -> str:
    ext = file_path.split(".")[-1].lower()
    try:
        if ext == "pdf":
            doc = fitz.open(file_path)
            return "".join(page.get_text() for page in doc)
        elif ext == "docx":
            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs)
        elif ext == "txt":
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
    except Exception as e:
        logging.exception(f"Extraction failed: {e}")
        return ""

def chunk_text_by_tokens(text: str, max_tokens: int = MAX_TOKENS_PER_CHUNK) -> list[str]:
    encoder = tiktoken.get_encoding("gpt2")
    tokens = encoder.encode(text)
    return [encoder.decode(tokens[i:i+max_tokens]) for i in range(0, len(tokens), max_tokens)]
    
def summarize_text_local(text: str, style: str = "bullet-point", model: str = MODEL_NAME) -> str:
    prompt_map = {
        "bullet-point": "Create a detailed, point-wise structured analysis of the following research paper including objectives, methods, contributions, challenges, and future work:",
        "paragraph": "Summarize the following research paper in a paragraph with key insights:",
        "detailed": "Write a comprehensive and structured summary of the following paper, broken into sections like Objective, Contributions, Challenges, Gaps, Future Work:"
    }
    prompt = f"{prompt_map.get(style, prompt_map['bullet-point'])}\n\n{text}"
    try:
        result = subprocess.run(
            [OLLAMA_PATH, "run", model],
            input=prompt,
            capture_output=True,
            text=True,
            timeout=180
        )
        if result.returncode == 0:
            return result.stdout.strip()
        else:
            logging.warning(f"Ollama error: {result.stderr}")
            return None
    except Exception as e:
        logging.exception(f"Ollama subprocess failed: {e}")
        return None

def generate_pdf(text: str, output_path: str, title: str = "Summary"):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=14)
    pdf.cell(0, 10, title, ln=True, align="C")
    pdf.ln(10)
    pdf.set_font("Arial", size=12)
    for line in text.split("\n"):
        pdf.multi_cell(0, 10, line)
    pdf.output(output_path)

# ------------------- Main Flow -------------------
uploaded_files = st.file_uploader("Upload PDF, DOCX, or TXT files", type=["pdf", "docx", "txt"], accept_multiple_files=True)

if uploaded_files:
    for uploaded_file in uploaded_files:
        filename = uploaded_file.name
        st.markdown(f"### 📄 Processing: **{filename}**")

        file_path = os.path.join(EXTRACTED_DIR, filename)
        with open(file_path, "wb") as f:
            f.write(uploaded_file.getbuffer())

        with st.expander("🔎 What happens during 'Analyzing'"):
            st.markdown("""
            1. ✅ Extracting plain text from PDF/DOCX/TXT.
            2. ✅ Cleaning and normalizing text (removing non-ASCII symbols).
            3. ✅ Tokenizing content using GPT-2 tokenizer.
            4. ✅ Splitting into ~800 token-sized chunks.
            5. ✅ Sending each chunk to local Ollama model (`mistral`) for summarization.
            6. ✅ Aggregating chunk-wise summaries.
            7. ✅ Generating final structured summary from all summaries.
            8. ✅ Saving the output summary as a PDF.
            """)

        with st.spinner("Extracting and chunking text..."):
            text = clean_text(extract_text_from_file(file_path))
            chunks = chunk_text_by_tokens(text)

        st.success(f"Split into {len(chunks)} chunks")
        partial_summaries = []
        chunk_data = []
        chunk_progress = st.progress(0, text="Summarizing chunks...")

        for i, chunk in enumerate(chunks):
            summary = summarize_text_local(chunk, style=depth)
            if not summary:
                summary = "[Error generating summary]"
            partial_summaries.append(summary)
            chunk_data.append({"Chunk #": i+1, "Summary": summary})
            chunk_progress.progress((i+1)/len(chunks), text=f"Chunk {i+1}/{len(chunks)}")

        chunk_progress.empty()

        if show_chunks:
            st.markdown("#### 🧩 Chunk Summaries")
            df = pd.DataFrame(chunk_data)
            gb = GridOptionsBuilder.from_dataframe(df)
            gb.configure_pagination()
            gb.configure_default_column(wrapText=True, autoHeight=True)
            AgGrid(df, gridOptions=gb.build(), theme="streamlit")

        combined = "\n".join(partial_summaries)
        final_summary = summarize_text_local(combined, style=depth)
        final_summary = final_summary or combined

        st.markdown("#### 🧾 Final Structured Summary")
        st.text_area("Summary:", value=final_summary, height=400)

        pdf_path = os.path.join(PDF_DIR, filename.rsplit(".", 1)[0] + "_summary.pdf")
        generate_pdf(final_summary, pdf_path)

        with open(pdf_path, "rb") as f:
            st.download_button("📥 Download Summary PDF", f, file_name=os.path.basename(pdf_path))

if show_logs:
    st.markdown("### 🐞 Debug Logs")
    if os.path.exists(LOG_FILE):
        with open(LOG_FILE, "r") as log:
            st.text(log.read())
